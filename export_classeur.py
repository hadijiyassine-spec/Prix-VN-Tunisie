# Export du classeur de prix VN vers un JSON de lignes, pour la fusion (integrer_prix.js).
#   python export_classeur.py sources/prix_VN_2011-092026_1.xlsx sources/classeur.json
#
# Ne fait AUCUNE interprétation métier : il lit, nettoie le format, et rend les lignes telles
# quelles. Les pièges traités ici sont uniquement des pièges de FORMAT :
#   · les feuilles 2026 et 2025 ont une ligne d'en-tête, les autres non, et 2011 commence par
#     une ligne vide — on reconnaît l'en-tête à son contenu, jamais à sa position ;
#   · la marque est collée au modèle par une espace INSÉCABLE jusqu'en 2025, par une espace
#     ordinaire en 2026 — on rend les deux formes, le rattachement se fait à la fusion ;
#   · cinq cellules Version sont des nombres (4, 160…) : converties en texte ;
#   · prix et variations sont du texte (« 149 990 DT », « -10 000 DT ») : convertis en entiers.
import sys, json, re
import openpyxl

src = sys.argv[1] if len(sys.argv) > 1 else 'sources/prix_VN_2011-092026_1.xlsx'
dst = sys.argv[2] if len(sys.argv) > 2 else 'sources/classeur.json'

EN_TETES = {'model', 'version', 'variation', 'prix', 'date', 'dealer'}
ESPACES = '\xa0    '


def texte(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)          # « 160.0 » → « 160 »
    s = str(v)
    for e in ESPACES:
        s = s.replace(e, ' ')
    return re.sub(r'\s+', ' ', s).strip()


def brut(v):
    """Texte SANS remplacer l'espace insécable : elle sépare la marque du modèle."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def nombre(v):
    """« 149 990 DT » → 149990 ; « -10 000 DT » → -10000 ; vide → None."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v)
    for e in ESPACES:
        s = s.replace(e, ' ')
    s = s.replace('DT', '').replace('dt', '').replace(' ', '').replace(',', '.')
    neg = s.startswith('-')
    s = re.sub(r'[^0-9.]', '', s)
    if not s:
        return None
    try:
        n = int(round(float(s)))
    except ValueError:
        return None
    return -n if neg else n


wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
lignes, ignorees = [], 0
for nom in wb.sheetnames:
    ws = wb[nom]
    for r in ws.iter_rows(values_only=True):
        if r is None:
            continue
        cells = list(r) + [None] * (6 - len(r))
        if all(c is None or str(c).strip() == '' for c in cells[:6]):
            continue
        bas = [texte(c).lower() for c in cells[:6]]
        if len([c for c in bas if c in EN_TETES]) >= 3:
            continue                                  # ligne d'en-tête
        model_brut = brut(cells[0])
        if not model_brut:
            ignorees += 1
            continue
        date = texte(cells[4])
        if not re.fullmatch(r'\d{2}\.\d{2}\.\d{4}', date):
            ignorees += 1
            continue
        lignes.append({
            'feuille': nom,
            'model_brut': model_brut,            # marque + modèle, séparateur d'origine
            'model': texte(cells[0]),            # espaces normalisées
            'version': texte(cells[1]),
            'variation': nombre(cells[2]),
            'prix': nombre(cells[3]),
            'date': date,
            'dealer': texte(cells[5]),
            'nbsp': '\xa0' in model_brut,
        })

with open(dst, 'w', encoding='utf-8') as f:
    json.dump(lignes, f, ensure_ascii=False)

print('%d lignes exportées vers %s (%d ignorées : modèle ou date absents)' % (len(lignes), dst, ignorees))
sansPrix = sum(1 for l in lignes if l['prix'] is None)
print('lignes sans prix exploitable :', sansPrix)
print('lignes avec espace insécable :', sum(1 for l in lignes if l['nbsp']))
